import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

output_dir = r"E:\EAUT\Kiểm thử phần mềm\TestCase"
os.makedirs(output_dir, exist_ok=True)

# Styles
font_title = Font(name="Arial", size=14, bold=True, color="1E293B")
font_subtitle = Font(name="Arial", size=9.5, italic=True, color="64748B")
font_meta = Font(name="Arial", size=9.5, bold=True, color="334155")

font_header = Font(name="Arial", size=9.5, bold=True, color="FFFFFF")
fill_header_unit = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
fill_header_int = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")  # Teal
fill_header_sys = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")  # Indigo

font_data = Font(name="Arial", size=9, color="0F172A")
font_bold_data = Font(name="Arial", size=9, bold=True, color="0F172A")

fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

fill_pass = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
font_pass = Font(name="Arial", size=9, bold=True, color="166534")

fill_fail = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
font_fail = Font(name="Arial", size=9, bold=True, color="991B1B")

fill_high = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
font_high = Font(name="Arial", size=9, bold=True, color="991B1B")

fill_med = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
font_med = Font(name="Arial", size=9, bold=True, color="92400E")

fill_low = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
font_low = Font(name="Arial", size=9, bold=True, color="075985")

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

def apply_sheet_formatting(ws, title, subtitle, header_fill, columns, data):
    last_col_letter = get_column_letter(len(columns))
    
    ws.merge_cells(f"A1:{last_col_letter}1")
    ws["A1"] = title
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells(f"A2:{last_col_letter}2")
    ws["A2"] = subtitle
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 18

    ws.merge_cells(f"A3:{last_col_letter}3")
    ws["A3"] = f"Hệ thống: Quản Lý Cửa Hàng Chăm Sóc Thú Cưng (PetCare Pro) | Tổng số: {len(data)} Test Cases | Trạng thái: 100% Đã Đối Soát Kết Quả Thực Tế"
    ws["A3"].font = font_meta
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 20

    ws.row_dimensions[4].height = 8

    # Header Row 5
    ws.row_dimensions[5].height = 26
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=5, column=col_idx, value=col_name)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border

    # Data Rows starting at 6
    for row_idx, row_data in enumerate(data, start=6):
        ws.row_dimensions[row_idx].height = 34
        is_even = (row_idx % 2 == 0)
        default_fill = fill_zebra if is_even else fill_white

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = font_data
            cell.fill = default_fill
            cell.border = thin_border

            col_header = columns[col_idx - 1]
            if col_header in ["STT", "Mã Test Case", "Mức Độ Ưu Tiên", "Mức Độ", "Trạng Thái", "Module", "Method / Hàm"]:
                cell.alignment = align_center
                if col_header == "Mã Test Case":
                    cell.font = font_bold_data
            else:
                cell.alignment = align_left

            val_str = str(value).upper()
            if val_str in ["HIGH", "CRITICAL", "CAO"]:
                cell.fill = fill_high
                cell.font = font_high
            elif val_str in ["MEDIUM", "TRUNG BÌNH"]:
                cell.fill = fill_med
                cell.font = font_med
            elif val_str in ["LOW", "THẤP"]:
                cell.fill = fill_low
                cell.font = font_low
            elif val_str in ["PASS", "ĐẠT"]:
                cell.fill = fill_pass
                cell.font = font_pass
            elif val_str in ["FAIL", "LỖI"]:
                cell.fill = fill_fail
                cell.font = font_fail

    # Auto set column widths
    for col_idx in range(1, len(columns) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row in range(5, 6 + len(data)):
            val = ws.cell(row=row, column=col_idx).value
            if val:
                lines = str(val).split("\n")
                line_max = max(len(l) for l in lines)
                if line_max > max_len:
                    max_len = line_max
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    ws.freeze_panes = "A6"

def add_members_sheet(wb, assigned_role):
    ws_m = wb.create_sheet(title="Thành Viên")
    ws_m.row_dimensions[1].height = 25
    headers = ["TT", "Họ Tên", "Mã Sinh Viên", "Lớp Hành Chính", "Phân Công Nhiệm Vụ"]
    for i, h in enumerate(headers, 1):
        c = ws_m.cell(1, i, h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        c.alignment = align_center
        c.border = thin_border

    members = [
        [1, "Đặng Văn Mạnh", "20232652", "DCCNTT14.C.3", "Trưởng nhóm - Xây dựng kịch bản Unit Test & Automation Test"],
        [2, "Trần Ngọc Sơn", "20232472", "DCCNTT14.C.3", "Thành viên - Xây dựng kịch bản Integration Test & API Test"],
        [3, "Ngô Hoàng Anh", "20232558", "DCCNTT14.C.3", "Thành viên - Xây dựng kịch bản System Test E2E & Thực thi"]
    ]

    for r_idx, m in enumerate(members, 2):
        ws_m.row_dimensions[r_idx].height = 24
        for c_idx, val in enumerate(m, 1):
            c = ws_m.cell(r_idx, c_idx, val)
            c.font = font_data
            c.border = thin_border
            c.alignment = align_center if c_idx in [1, 3, 4] else align_left
            if c_idx == 5 and assigned_role in val:
                c.font = font_bold_data
                c.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

    for col_idx in range(1, 6):
        ws_m.column_dimensions[get_column_letter(col_idx)].width = 25

print("Helper functions defined.")
